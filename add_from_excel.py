"""Resolve invite links from an xlsx, join if needed, add new groups to the DB."""

from __future__ import annotations

import asyncio
import logging
import random
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

sys.path.insert(0, str(Path(__file__).resolve().parent))

from config import get_settings, setup_logging  # noqa: E402
from db import Group, init_db, session_scope  # noqa: E402
from join_pending import MAX_DELAY_SECONDS, MIN_DELAY_SECONDS, join_one  # noqa: E402
from web_transport import WebTransport  # noqa: E402

logger = logging.getLogger("eitaa.add_xlsx")
JOIN = re.compile(r"joinchat\.?/([A-Za-z0-9_+-]+)")


def hashes_from_xlsx(path: Path) -> list[str]:
    wb = load_workbook(path, data_only=True)
    seen: set[str] = set()
    hashes: list[str] = []
    for row in wb.active.iter_rows(values_only=True):
        value = row[1] if row and len(row) > 1 else None
        if not value:
            continue
        match = JOIN.search(str(value))
        if not match:
            logger.warning("unparsed row: %s", value)
            continue
        hash_ = match.group(1)
        if hash_ not in seen:
            seen.add(hash_)
            hashes.append(hash_)
    return hashes


async def check_invite(transport: WebTransport, hash_: str) -> dict:
    driver = transport._require_driver()  # noqa: SLF001

    def _check() -> dict:
        script = f"""
        return apiManager.invokeApi('messages.checkChatInvite', {{hash: '{hash_}'}})
          .then(r => {{
            const chat = r.chat || (r.chats && r.chats[0]) || null;
            return {{
              ok: true,
              type: r._ || null,
              already: r._ === 'chatInviteAlready' || !!(chat && chat.id),
              id: chat ? chat.id : null,
              title: (chat && chat.title) || r.title || null
            }};
          }})
          .catch(e => ({{
            ok: false,
            error: (e && (e.text || e.type || e.message)) || String(e)
          }}));
        """
        return driver.execute_script(script)

    return await transport._run(_check)  # noqa: SLF001


async def upsert_group(chat_id: str, title: str | None) -> str:
    async with session_scope() as session:
        existing = (
            await session.execute(select(Group).where(Group.chat_id == chat_id))
        ).scalar_one_or_none()
        if existing is None:
            session.add(Group(chat_id=chat_id, title=title))
            return "added"
        if title and not existing.title:
            existing.title = title
        if not existing.enabled and existing.status != "kicked":
            existing.enabled = True
            existing.status = "active"
            existing.fail_count = 0
            existing.last_error = None
            return "re-enabled"
        return "exists"


async def main() -> int:
    setup_logging()
    await init_db()
    xlsx = Path(sys.argv[1] if len(sys.argv) > 1 else r"c:\Users\MR\Downloads\eitaa (2).xlsx")
    hashes = hashes_from_xlsx(xlsx)
    logger.info("%d unique invite hashes in %s", len(hashes), xlsx.name)

    async with session_scope() as session:
        existing_ids = set(
            (await session.execute(select(Group.chat_id))).scalars().all()
        )

    transport = WebTransport(get_settings())
    await transport.start(headless=False)
    added = joined = skipped = failed = 0
    try:
        if not await transport.is_logged_in():
            logger.error("not logged in")
            return 1

        pending: list[tuple[str, str | None]] = []
        for index, hash_ in enumerate(hashes, start=1):
            info = await check_invite(transport, hash_)
            if not info.get("ok"):
                failed += 1
                logger.warning("[%d/%d] resolve failed %s: %s", index, len(hashes), hash_, info.get("error"))
                continue
            title = info.get("title")
            chat_id = str(-info["id"]) if info.get("id") else None
            if chat_id and chat_id in existing_ids:
                skipped += 1
                logger.info("[%d/%d] already in db: %s %s", index, len(hashes), chat_id, title)
                continue
            if info.get("already") and chat_id:
                action = await upsert_group(chat_id, title)
                existing_ids.add(chat_id)
                added += 1
                logger.info("[%d/%d] %s (already member): %s %s", index, len(hashes), action, chat_id, title)
                continue
            pending.append((hash_, title))
            logger.info("[%d/%d] need join: %s", index, len(hashes), title or hash_)

        logger.info("%d groups to join", len(pending))
        for index, (hash_, title) in enumerate(pending, start=1):
            result = await join_one(transport, hash_)
            if result.get("ok") and result.get("id"):
                chat_id = str(-result["id"])
                action = await upsert_group(chat_id, result.get("title") or title)
                existing_ids.add(chat_id)
                joined += 1
                added += 1
                logger.info("[%d/%d] %s after join: %s %s", index, len(pending), action, chat_id, result.get("title") or title)
            else:
                failed += 1
                logger.warning("[%d/%d] join failed %s: %s", index, len(pending), title or hash_, result.get("error"))
            if index < len(pending):
                delay = random.uniform(MIN_DELAY_SECONDS, MAX_DELAY_SECONDS)
                logger.info("waiting %.0fs", delay)
                await asyncio.sleep(delay)
    finally:
        await transport.aclose()

    logger.info("done added=%s joined=%s skipped=%s failed=%s", added, joined, skipped, failed)
    return 0 if failed == 0 else 1


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
